# -*- coding: utf-8 -*-
"""
Import ຕາຕາລາງເວນຈາກ Excel ເຂົ້າ Supabase
Usage:
  pip install openpyxl supabase python-dotenv
  python scripts/import-from-xlsx.py "d:\path\to\file.xlsx" --year 2026 --month 6
"""
import argparse
import os
import re
import sys
from datetime import date
from pathlib import Path

from dotenv import load_dotenv

ROOT = Path(__file__).resolve().parents[1]
load_dotenv(ROOT / '.env')

try:
    import openpyxl
    from supabase import create_client
except ImportError:
    print('pip install openpyxl supabase python-dotenv')
    sys.exit(1)


def parse_sheet(ws, year: int, month: int):
    """อ่านแถวพยาบาล + เวน จาก Sheet2 layout Luckxay"""
    staff_rows = []
    header_day_col = None

    for r in range(1, ws.max_row + 1):
        c1 = ws.cell(r, 1).value
        c2 = ws.cell(r, 2).value
        if c2 and 'ຊື່' in str(c2):
            header_day_col = 3
            continue
        if header_day_col and isinstance(c1, (int, float)) and c2 and isinstance(c2, str):
            name = str(c2).strip()
            shifts = {}
            for day in range(1, 32):
                col = header_day_col + day - 1
                val = ws.cell(r, col).value
                if val is None or str(val).strip() == '':
                    continue
                code = str(val).strip().upper()
                if code == 'B/S':
                    code = 'B/S'
                shifts[day] = code
            staff_rows.append({'name': name, 'shifts': shifts})

    return staff_rows


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('xlsx', help='path to Excel file')
    ap.add_argument('--year', type=int, default=2026)
    ap.add_argument('--month', type=int, default=6)
    ap.add_argument('--dry-run', action='store_true')
    args = ap.parse_args()

    url = os.environ.get('SUPABASE_URL')
    key = os.environ.get('SUPABASE_SERVICE_KEY')
    if not url or not key or len(key) < 80:
        print('ERROR: ໃສ່ SUPABASE_URL ແລະ SUPABASE_SERVICE_KEY ທັງໝົດໃນ .env')
        sys.exit(1)

    wb = openpyxl.load_workbook(args.xlsx, data_only=True)
    ws = wb.active
    staff_rows = parse_sheet(ws, args.year, args.month)
    print(f'Found {len(staff_rows)} staff')

    if args.dry_run:
        for s in staff_rows[:3]:
            print(s['name'], len(s['shifts']), 'days')
        return

    sb = create_client(url, key)

    # ลบเวนเดือนนี้ก่อน (optional — comment out if merge only)
    # sb.table('schedules').delete().eq('year', args.year).eq('month', args.month).execute()

    name_to_id = {}
    existing = sb.table('staff').select('id,name').eq('is_active', True).execute()
    for row in existing.data or []:
        name_to_id[row['name']] = row['id']

    sort_order = len(name_to_id)
    inserts = []
    for s in staff_rows:
        if s['name'] not in name_to_id:
            sort_order += 1
            res = sb.table('staff').insert({
                'name': s['name'],
                'position': 'Nurse',
                'department': 'Nursing',
                'sort_order': sort_order,
            }).execute()
            name_to_id[s['name']] = res.data[0]['id']
            print('+ staff:', s['name'])

        staff_id = name_to_id[s['name']]
        for day, code in s['shifts'].items():
            try:
                d = date(args.year, args.month, day)
            except ValueError:
                continue
            work_date = d.isoformat()
            # normalize BS
            if code == 'BS':
                code = 'B/S'
            inserts.append({
                'staff_id': staff_id,
                'shift_code': code,
                'work_date': work_date,
            })

    # batch upsert
    batch_size = 100
    for i in range(0, len(inserts), batch_size):
        chunk = inserts[i:i + batch_size]
        sb.table('schedules').upsert(chunk, on_conflict='staff_id,work_date').execute()
        print(f'  upserted {i + len(chunk)}/{len(inserts)}')

    print('Done.')


if __name__ == '__main__':
    main()
