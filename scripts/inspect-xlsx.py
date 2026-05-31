# -*- coding: utf-8 -*-
import sys
import openpyxl

path = r'd:\Luckxay_Hospital\ID\ອັບເດດຕາຕາລາງເດືອນ6.xlsx'
wb = openpyxl.load_workbook(path, data_only=True)
sys.stdout.reconfigure(encoding='utf-8')

for sn in wb.sheetnames:
    ws = wb[sn]
    print(f'=== {sn} {ws.max_row}x{ws.max_column} ===')
    for r in range(1, min(ws.max_row + 1, 70)):
        cells = []
        for c in range(1, min(ws.max_column + 1, 50)):
            v = ws.cell(r, c).value
            if v is not None and str(v).strip():
                cells.append((c, v))
        if cells:
            print(f'R{r}:', cells[:40])
